"""Build the dataspace knowledge graph from the AM2Scale workbook.

Reads data/AM2Scale_Mini_Datensatz_erweitert.xlsx, maps every sheet onto the
canonical KG entities from kg-schema.json, and publishes one subgraph per
dataspace participant to the discovery service — where the free-text fields
become semantically searchable and the structured ones become filters.

Who owns what follows the workbook: printers (and everything produced on them)
belong to the company in Unternehmens_ID; the ERP materials carry no owner in
the data, so they are distributed across the participants for the demo
(--material-owner pins them to a single company instead).

Materials are additionally enriched from data/material-semantik.json (or a
"Semantik" sheet in the workbook, which takes precedence) — Einsatzzweck,
Funktion, Anforderungen, Branche, Synonyme. Without that free text a semantic
search over part master data has very little to work with.

    python import_dataset.py --all
    python import_dataset.py --all --dry-run --out graph.json
    python import_dataset.py --all --material-owner fha-wien
"""
import argparse
import json
import re
import statistics
import sys
import urllib.error
import urllib.request
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent
REPO = ROOT.parent.parent
DATASET = REPO / "data" / "AM2Scale_Mini_Datensatz_erweitert.xlsx"
SEMANTIK = REPO / "data" / "material-semantik.json"
SCHEMA = json.loads((ROOT / "kg-schema.json").read_text(encoding="utf-8"))
DISCOVERY_URL = "http://127.0.0.1:5185"
TRAEFIK = "http://127.0.0.1:80"
API_KEY = "password"
COMPANIES_DIR = REPO / "compose" / "companies"

# dataspace participant -> dataset company (U_ID) it plays
DEFAULT_MAPPING = {
    "fha-wien":          "U6",   # Fraunhofer Austria Wien — 3 Drucker
    "fha-tirol":         "U7",   # Fraunhofer Austria — 1 Drucker (s. COMPANY_OVERRIDES)
    "oebb":              "U1",   # ÖBB Wien
    "wiener-linien":     "U2",   # WienerLinien
    "wien-energie":      "U3",   # WienEnergie
    "wiener-netze":      "U4",   # WienerNetze
    "wiener-stadtwerke": "U0",   # Konzernholding — synthetisch, s. SYNTHETIC_FIRMS
}

# U5 (ÖBB St. Pölten) is deliberately unused: it owns no printers, so it would
# add nothing but a second, dataless ÖBB participant.

# Companies the dataset does not contain. Same shape as an Unternehmensdaten row.
SYNTHETIC_FIRMS = {
    "U0": {"U_ID": "U0", "Unternehmensname": "Wiener Stadtwerke",
           "Straße": "Thomas-Klestil-Platz 14", "Stadt": "Wien", "Land": "Austria",
           "Stromtarif": 0.06, "CO2_KWh": 0.09},
}

# Corrections applied on top of the Unternehmensdaten row, keyed by sheet column.
# Lets a participant present its real identity without touching the Excel file.
# fha-tirol plays U7, which the dataset records as the Graz site.
COMPANY_OVERRIDES = {
    "fha-wien":  {"Unternehmensname": "Fraunhofer Austria Wien"},
    "fha-tirol": {"Unternehmensname": "Fraunhofer Austria Tirol",
                  "Stadt": "Innsbruck", "Straße": "Karl-Kapferer-Straße 5"},
    "oebb":      {"Unternehmensname": "ÖBB"},
}

# The ERP short texts are raw codes (PYC3D_DUCK_KEYCHAIN). These give them a
# readable name for the UI; the code itself stays in the asset name so the
# record remains traceable back to the ERP. Unlisted materials fall back to a
# tidied-up short text.
MATERIAL_LABELS = {
    "Worldcup_386":         "Pokal",
    "PYC3D_DUCK_KEYCHAIN":  "Enten-Schlüsselanhänger",
    "Adapterplatte_47":     "Adapterplatte",
    "Gehaeusedeckel_Klein": "Gehäusedeckel klein",
}


def material_label(kurztext):
    """Readable label for an ERP short text, or None if there is nothing to show."""
    if not kurztext:
        return None
    raw = str(kurztext).strip()
    known = MATERIAL_LABELS.get(raw)
    if known:
        return known
    # drop a trailing running number, turn separators into spaces
    tidied = re.sub(r"[_-]+\d+$", "", raw)
    tidied = re.sub(r"[_-]+", " ", tidied).strip()
    return tidied or None


SEMANTIC_FIELDS = ("einsatzzweck", "funktion", "anforderungen", "branche", "synonyme")


# --------------------------------------------------------------- excel utils

def sheet_dicts(wb, name):
    """Rows of a sheet as dicts keyed by header, skipping fully empty rows."""
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else None for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or all(v is None for v in r):
            continue
        out.append({h: v for h, v in zip(header, r) if h})
    return out


def clean(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def map_attrs(entity_name, row):
    """Map a sheet row onto the entity's schema attributes."""
    spec = SCHEMA["entities"][entity_name]["attributes"]
    attrs = {}
    for attr, meta in spec.items():
        src = meta.get("from")
        if not src:
            continue
        val = clean(row.get(src))
        if val is not None:
            attrs[attr] = val
    return attrs


def load_semantics(wb):
    """Semantic enrichment per material number: workbook sheet wins over file."""
    sem = {}
    if SEMANTIK.exists():
        try:
            data = json.loads(SEMANTIK.read_text(encoding="utf-8"))
            for matnr, fields in (data.get("materials") or {}).items():
                sem[str(matnr).strip()] = {k: v for k, v in fields.items()
                                           if k in SEMANTIC_FIELDS and v}
        except (json.JSONDecodeError, OSError) as e:
            print(f"  ! material-semantik.json nicht lesbar ({e})")
    for row in sheet_dicts(wb, "Semantik"):
        matnr = str(row.get("Materialnummer") or row.get("Material") or "").strip()
        if not matnr:
            continue
        entry = sem.setdefault(matnr, {})
        for f in SEMANTIC_FIELDS:
            for key in (f, f.capitalize()):
                if row.get(key):
                    entry[f] = clean(row[key])
    return sem


# ------------------------------------------------------------- graph builder

def node(node_id, ntype, attrs):
    return {"id": node_id, "type": ntype, "attrs": attrs}


def minutes_between(a, b):
    try:
        return round((datetime.fromisoformat(str(b)) - datetime.fromisoformat(str(a))).total_seconds() / 60, 1)
    except (ValueError, TypeError):
        return None


def aggregate_machine_data(wb, printer_serial, include_metrics=False):
    """Condense the MTConnect series of one printer into per-job nodes."""
    rows = [r for r in sheet_dicts(wb, "Maschinenzeitreihendaten")
            if str(r.get("printerSerial") or "").strip() == printer_serial]
    if not rows:
        return []
    by_job = {}
    for r in rows:
        by_job.setdefault(str(r.get("pathProgram") or "unbekannt").strip(), []).append(r)

    out = []
    for job, jrows in by_job.items():
        stamps = sorted(str(clean(r.get("_timestamp_log"))) for r in jrows if r.get("_timestamp_log"))
        attrs = {
            "serieId": f"{printer_serial}:{job}",
            "printerId": printer_serial,
            "jobProgramm": job,
            "messwerte": len(jrows),
            "zeitraumVon": stamps[0] if stamps else None,
            "zeitraumBis": stamps[-1] if stamps else None,
            "rohdatenRef": f"maschinendaten://{printer_serial}/{job}",
        }
        if include_metrics:
            kennzahlen = {}
            for ch in ("m1TAct", "oven1CurrentAct", "z1Act", "curLayer"):
                vals = [float(r[ch]) for r in jrows if isinstance(r.get(ch), (int, float))]
                if len(vals) >= 3:
                    kennzahlen[ch] = {"min": round(min(vals), 3), "max": round(max(vals), 3),
                                      "avg": round(statistics.fmean(vals), 3)}
            if kennzahlen:
                attrs["kennzahlen"] = kennzahlen
        out.append(node(f"serie:{printer_serial}:{job}", "Maschinendaten",
                        {k: v for k, v in attrs.items() if v is not None}))
    return out


def build_graphs(wb, mapping, material_owner=None, include_metrics=False):
    """Build every participant's subgraph in one pass, so cross-references
    (printer -> company, material -> orders) stay consistent."""
    semantics = load_semantics(wb)

    firms = {str(r.get("U_ID")).strip(): r for r in sheet_dicts(wb, "Unternehmensdaten")}
    for uid, row in SYNTHETIC_FIRMS.items():
        firms.setdefault(uid, row)
    uid_to_company = {uid: comp for comp, uid in mapping.items()}

    graphs = {}
    for company, uid in mapping.items():
        firm_row = firms.get(uid)
        if not firm_row:
            raise SystemExit(f"U_ID '{uid}' nicht im Dataset (vorhanden: {sorted(firms)})")
        override = COMPANY_OVERRIDES.get(company)
        if override:
            firm_row = {**firm_row, **override}
        graphs[company] = {
            "owner": company,
            "did": f"did:web:identityhub-{company}%3A7083:{company}",
            "datasetCompany": {"uId": uid, "name": clean(firm_row.get("Unternehmensname")),
                               "stadt": clean(firm_row.get("Stadt"))},
            "nodes": [node(f"unternehmen:{uid}", "Unternehmen", map_attrs("Unternehmen", firm_row))],
            "edges": [],
        }

    def add(company, n):
        graphs[company]["nodes"].append(n)

    def link(company, a, rel, b):
        graphs[company]["edges"].append({"from": a, "rel": rel, "to": b})

    # --- Drucker: owned by the company in Unternehmens_ID -------------------
    printer_owner, printer_serial, printer_model = {}, {}, {}
    for r in sheet_dicts(wb, "Druckerdaten"):
        pid = str(r.get("PrinterID") or "").strip()
        if not pid:
            continue
        printer_model[pid] = str(r.get("Printer_Model") or "").strip()
        owner = uid_to_company.get(str(r.get("Unternehmens_ID") or "").strip())
        if not owner:
            continue  # printer of a company that is not a dataspace participant
        printer_owner[pid] = owner
        printer_serial[pid] = str(r.get("Seriennummer") or "").strip()
        add(owner, node(f"drucker:{pid}", "Drucker", map_attrs("Drucker", r)))
        link(owner, f"unternehmen:{mapping[owner]}", "betreibt", f"drucker:{pid}")

    # --- Material: no owner in the data -> distribute (or pin) --------------
    materials = sheet_dicts(wb, "ERP")
    participants = list(mapping.keys())
    material_owner_of = {}
    for i, r in enumerate(materials):
        matnr = str(r.get("Material") or "").strip()
        if not matnr:
            continue
        owner = material_owner or participants[i % len(participants)]
        material_owner_of[matnr] = owner
        attrs = map_attrs("Material", r)
        attrs.update(semantics.get(matnr, {}))
        # readable title for the UIs; the raw short text stays alongside it
        label = material_label(attrs.get("materialkurztext"))
        if label and label != str(attrs.get("materialkurztext")):
            attrs["benennung"] = label
        add(owner, node(f"material:{matnr}", "Material", attrs))
        link(owner, f"unternehmen:{mapping[owner]}", "fuehrt", f"material:{matnr}")

    def owner_of_material(matnr):
        return material_owner_of.get(str(matnr).strip())

    # --- Druckprofil: lives with the material's owner ------------------------
    # A print profile is knowledge about the part ("this can be made on a
    # Prusa MK4S"), so it belongs to whoever offers the part — even when the
    # machine itself is operated by someone else. The printer model is carried
    # along so the filter works across company boundaries; the edge to the
    # actual printer node is only drawn when both are in the same subgraph.
    for r in sheet_dicts(wb, "Bauteildruckdaten"):
        pid = str(r.get("PrinterID") or "").strip()
        matnr = str(r.get("Materialnummer") or "").strip()
        owner = owner_of_material(matnr)
        if not owner:
            continue
        attrs = map_attrs("Druckprofil", r)
        attrs["profilId"] = f"{matnr}_{pid}"
        if pid in printer_model:
            attrs["druckerModell"] = printer_model[pid]
        nid = f"profil:{attrs['profilId']}"
        add(owner, node(nid, "Druckprofil", attrs))
        link(owner, nid, "fuer", f"material:{matnr}")
        if printer_owner.get(pid) == owner:
            link(owner, nid, "laeuft_auf", f"drucker:{pid}")

    # --- Auftragsposition: with the material owner (they order it) ----------
    for r in sheet_dicts(wb, "Auftragsdaten"):
        matnr = str(r.get("Materialnummer") or "").strip()
        owner = owner_of_material(matnr)
        if not owner:
            continue
        attrs = map_attrs("Auftragsposition", r)
        attrs["positionId"] = f"{attrs.get('auftragsnummer')}-{attrs.get('position')}"
        nid = f"position:{attrs['positionId']}"
        add(owner, node(nid, "Auftragsposition", attrs))
        link(owner, nid, "bestellt", f"material:{matnr}")

    # --- Fertigungsauftrag: with the printer's owner ------------------------
    charge_owner, charge_of_fa = {}, {}
    for r in sheet_dicts(wb, "Fertigungsauftragsdaten"):
        fa = str(r.get("FertigungsauftragsID") or "").strip()
        pid = str(r.get("PrinterID") or "").strip()
        owner = printer_owner.get(pid)
        if not fa or not owner:
            continue
        attrs = map_attrs("Fertigungsauftrag", r)
        # charge id follows the pattern used in Bauteildatenbank/Qualitätsdaten: A001-F001
        auf = str(attrs.get("auftragsnummer") or "").replace("-", "")
        charge = f"{auf}-{fa.replace('-', '')}" if auf else fa
        attrs["chargenId"] = charge
        dauer = minutes_between(attrs.get("startzeitpunkt"), attrs.get("endzeitpunkt"))
        if dauer is not None:
            attrs["dauerMin"] = dauer
        nid = f"fertigung:{fa}"
        add(owner, node(nid, "Fertigungsauftrag", attrs))
        link(owner, nid, "laeuft_auf", f"drucker:{pid}")
        charge_owner[charge] = owner
        charge_of_fa[charge] = nid
        matnr = str(attrs.get("materialnummer") or "").strip()
        if owner_of_material(matnr) == owner:
            link(owner, nid, "fertigt", f"material:{matnr}")
        pos_id = f"{attrs.get('auftragsnummer')}-{attrs.get('position')}"
        if any(n["id"] == f"position:{pos_id}" for n in graphs[owner]["nodes"]):
            link(owner, nid, "erfuellt", f"position:{pos_id}")

    # --- Bauteil + DPP-Ereignisse -------------------------------------------
    dpp_by_part = {}
    for r in sheet_dicts(wb, "Digitaler-Produktpass"):
        pid = str(r.get("Bauteil-ID") or "").strip()
        if pid:
            dpp_by_part.setdefault(pid, []).append({
                "zeitstempel": clean(r.get("Zeitstempel")),
                "ereignis": clean(r.get("Ereignis/Änderung")),
                "bearbeiter": clean(r.get("Bearbeiter")),
                "version": clean(r.get("Neue DPP-Version")),
                "status": clean(r.get("Status")),
            })

    # Bauteil-IDs are only unique within a charge in this workbook (123456789_B1
    # appears in several charges), and some rows are duplicated outright — so the
    # node id combines charge + part id, and exact repeats are collapsed.
    seen_parts, dup_rows = set(), 0
    for r in sheet_dicts(wb, "Bauteildatenbank"):
        bid = str(r.get("Bauteil-ID") or "").strip()
        charge = str(r.get("ChargenID") or "").strip()
        owner = charge_owner.get(charge)
        if not bid or not owner:
            continue
        nid = f"bauteil:{charge}:{bid}"
        if nid in seen_parts:
            dup_rows += 1
            continue
        seen_parts.add(nid)
        attrs = map_attrs("Bauteil", r)
        events = dpp_by_part.get(bid, [])
        if events:
            attrs["ereignisse"] = events
            attrs["dppVersion"] = events[-1].get("version")
            attrs["dppStatus"] = events[-1].get("status")
        add(owner, node(nid, "Bauteil", attrs))
        if charge in charge_of_fa:
            link(owner, charge_of_fa[charge], "erzeugt", nid)
    if dup_rows:
        print(f"  i {dup_rows} doppelte Zeilen in 'Bauteildatenbank' übersprungen "
              f"(gleiche Bauteil-ID in derselben Charge)")

    # --- Qualitätsprüfung ----------------------------------------------------
    for r in sheet_dicts(wb, "Qualitätsdaten"):
        charge = str(r.get("Charge") or "").strip()
        owner = charge_owner.get(charge)
        if not owner:
            continue
        attrs = map_attrs("Qualitaetspruefung", r)
        attrs["pruefungId"] = f"QS-{charge}"
        nid = f"qs:{attrs['pruefungId']}"
        add(owner, node(nid, "Qualitaetspruefung", attrs))
        if charge in charge_of_fa:
            link(owner, charge_of_fa[charge], "geprueft_in", nid)

    # --- Maschinenzeitreihen -------------------------------------------------
    for pid, owner in printer_owner.items():
        serial = printer_serial.get(pid)
        if not serial:
            continue
        for n in aggregate_machine_data(wb, serial, include_metrics):
            n["attrs"]["printerId"] = pid
            add(owner, n)
            link(owner, f"drucker:{pid}", "liefert", n["id"])

    return graphs


# ------------------------------------------------- EDC assets for KG parts

def mgmt(company, method, path, payload=None, timeout=30):
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(f"{TRAEFIK}{path}", data=body, method=method)
    req.add_header("Host", f"cp.{company}.localhost")
    req.add_header("X-Api-Key", API_KEY)
    if body is not None:
        req.add_header("Content-Type", "application/json")
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
        return json.loads(raw.decode("utf-8")) if raw else None


def create_material_asset(company, mat_node):
    """Publish a material's master-data sheet as a retrievable EDC asset, so a
    search hit is not just metadata but something a partner can fetch."""
    attrs = mat_node["attrs"]
    matnr = attrs.get("materialnummer") or mat_node["id"].split(":")[-1]
    asset_id = f"stammdaten-{matnr}"
    filename = f"{asset_id}.json"

    spec = SCHEMA["entities"]["Material"]["attributes"]
    sheet = {k: v for k, v in attrs.items() if spec.get(k, {}).get("tier") == 1}
    store = COMPANIES_DIR / company / "storage" / "assets"
    store.mkdir(parents=True, exist_ok=True)
    (store / filename).write_text(json.dumps(sheet, ensure_ascii=False, indent=1), encoding="utf-8")

    kurztext = attrs.get("materialkurztext") or matnr
    label = material_label(attrs.get("materialkurztext"))
    display = (f"Stammdaten {label} ({kurztext})"
               if label and label != str(kurztext) else f"Stammdaten {kurztext}")

    properties = {
        "name": display,
        "description": attrs.get("einsatzzweck") or f"Stammdaten zu {matnr}",
        "am2scale:partName": attrs.get("materialkurztext") or matnr,
        "am2scale:material": attrs.get("materialkurztext", ""),
        "am2scale:werkstoff": attrs.get("werkstoff", ""),
        "am2scale:abmasse": " x ".join(str(attrs[k]) for k in ("laenge", "breite", "hoehe")
                                       if attrs.get(k) is not None) or "",
        "am2scale:kgNode": mat_node["id"],
        "am2scale:tier": "1",
        "am2scale:fileName": filename,
        "am2scale:fileFormat": "JSON",
    }
    try:
        for path, payload in (
            ("/api/mgmt/v4/assets", {
                "@context": ["https://w3id.org/edc/connector/management/v2"],
                "@id": asset_id, "@type": "Asset", "properties": properties,
                "dataAddress": {"@type": "DataAddress", "type": "HttpData",
                                "baseUrl": f"http://filestore-{company}/{filename}",
                                "proxyPath": "true", "proxyQueryParams": "true"}}),
            ("/api/mgmt/v4/contractdefinitions", {
                "@context": ["https://w3id.org/edc/connector/management/v2"],
                "@id": f"{asset_id}-def", "@type": "ContractDefinition",
                "accessPolicyId": f"{company}-require-membership",
                "contractPolicyId": f"{company}-require-manufacturer",
                "assetsSelector": {"@type": "Criterion",
                                   "operandLeft": "https://w3id.org/edc/v0.0.1/ns/id",
                                   "operator": "=", "operandRight": asset_id}}),
        ):
            try:
                mgmt(company, "POST", path, payload)
            except urllib.error.HTTPError as e:
                if e.code != 409:
                    raise
                # already there: update it, otherwise renames/edits never land
                if path.endswith("/assets"):
                    mgmt(company, "PUT", path, payload)
        return asset_id
    except (urllib.error.HTTPError, urllib.error.URLError) as e:
        print(f"  ! EDC-Asset für {matnr} nicht angelegt ({e}) — läuft der Stack?")
        return None


# ------------------------------------------------------------------ publish

def publish(graph):
    body = json.dumps(graph).encode("utf-8")
    req = urllib.request.Request(f"{DISCOVERY_URL}/kg/publish", data=body, method="POST")
    req.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise SystemExit(f"Discovery lehnte ab (HTTP {e.code}): {e.read().decode(errors='replace')[:300]}")
    except urllib.error.URLError as e:
        raise SystemExit(f"Discovery-Service nicht erreichbar unter {DISCOVERY_URL} ({e}). "
                         f"Zuerst .\\start-discovery.ps1 starten.")


def main():
    ap = argparse.ArgumentParser(description="Baut den Dataspace-KG aus dem AM2Scale-Datensatz")
    ap.add_argument("--all", action="store_true", help="alle Teilnehmer laut Standard-Mapping")
    ap.add_argument("--company", help="nur diesen Teilnehmer publizieren")
    ap.add_argument("--material-owner", help="alle Materialien dieser Firma zuordnen "
                                             "(Standard: gleichmäßig verteilt)")
    ap.add_argument("--no-assets", action="store_true", help="keine EDC-Assets anlegen")
    ap.add_argument("--include-metrics", action="store_true",
                    help="Kennzahlen aus den Maschinenzeitreihen aufnehmen")
    ap.add_argument("--dry-run", action="store_true", help="nur bauen, nicht publizieren")
    ap.add_argument("--out", help="Graphen zusätzlich als JSON schreiben")
    args = ap.parse_args()

    if not args.all and not args.company:
        ap.error("--all oder --company angeben")
    if not DATASET.exists():
        raise SystemExit(f"Dataset nicht gefunden: {DATASET}")

    print(f"Lade {DATASET.name} ...")
    wb = openpyxl.load_workbook(DATASET, read_only=True, data_only=True)
    graphs = build_graphs(wb, DEFAULT_MAPPING, args.material_owner, args.include_metrics)

    wanted = [args.company] if args.company else list(graphs)
    for company in wanted:
        g = graphs.get(company)
        if not g:
            raise SystemExit(f"Unbekannter Teilnehmer '{company}' (bekannt: {sorted(graphs)})")
        by_type = {}
        for n in g["nodes"]:
            by_type[n["type"]] = by_type.get(n["type"], 0) + 1
        ds = g["datasetCompany"]
        print(f"\n{company}  (spielt {ds['uId']} / {ds['name']}, {ds.get('stadt')})")
        print(f"  Knoten: {len(g['nodes']):>3}  {by_type}")
        print(f"  Kanten: {len(g['edges']):>3}")

        if not args.no_assets and not args.dry_run:
            for n in g["nodes"]:
                if n["type"] == "Material":
                    aid = create_material_asset(company, n)
                    if aid:
                        n["attrs"]["edcAssetId"] = aid
                        print(f"  EDC-Asset: {aid}  ({n['attrs'].get('materialkurztext')})")
        if args.out:
            out = Path(args.out)
            out = out.with_name(f"{out.stem}-{company}{out.suffix or '.json'}")
            out.write_text(json.dumps(g, ensure_ascii=False, indent=1), encoding="utf-8")
            print(f"  geschrieben: {out}")
        if not args.dry_run:
            print(f"  publiziert: {publish(g)}")

    if args.dry_run:
        print("\n(dry-run — nichts publiziert)")


if __name__ == "__main__":
    sys.exit(main())
